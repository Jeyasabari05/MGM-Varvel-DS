"""MGM Varvel motor datasheet selector. Source workbooks stay the authority."""
from __future__ import annotations

import json
import mimetypes
import re
import sys
from decimal import Decimal, InvalidOperation
from collections import defaultdict
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import unquote, urlparse

try:
    from openpyxl import load_workbook
    import pymupdf as fitz
except ImportError as exc:
    raise SystemExit(f"Missing dependency: {exc.name}. Install packages from requirements.txt")

ROOT = Path(__file__).resolve().parent
SOURCES = ROOT / "data" / "source_files"
PDF = ROOT / "templates" / "DS-025-2028.pdf"
STATIC = ROOT / "static"
FAMILIES = {"BAX": "BAX", "BMX": "BMX", "SMX": "SMX"}
SOURCE_CLASS = {
    "1. BA - IE3 - AC.xlsx": ("BAX", "IE3", "AC"),
    "2. BA - IE3 - DC.xlsx": ("BAX", "IE3", "DC"),
    "3. BM - IE3.xlsx": ("BMX", "IE3", None),
    "4. SM - IE3.xlsx": (None, "IE3", None),
    "5. BA - IE2 - AC.xlsx": ("BAX", "IE2", "AC"),
    "6. BA - IE2 - DC.xlsx": ("BAX", "IE2", "DC"),
    "7. BM - IE2.xlsx": ("BMX", "IE2", "DC"),
    "8. SM - IE2.xlsx": ("SMX", "IE2", None),
    "SM - IE2 (upto 55kW).xlsx": ("SMX", "IE2", None),
    "SM - IE3 (upto 55kW).xlsx": ("SMX", "IE3", None),
}

def clean(v):
    if v is None: return None
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    s = str(v).strip()
    return s if s else None

def excel_display(v, number_format="General"):
    """Respect the cell's displayed precision while retaining source units."""
    if v is None: return None
    if isinstance(v, (int, float)) and number_format and number_format != "General":
        decimals = re.search(r"\.([0#]+)", number_format)
        if decimals: return f"{v:.{len(decimals.group(1))}f}"
        if number_format.startswith("0") and "." not in number_format: return str(int(round(v)))
    return clean(v)

def read_sources():
    records, errors = [], []
    for name, (family_hint, class_hint, brake_hint) in SOURCE_CLASS.items():
        path = SOURCES / name
        if not path.exists():
            errors.append(f"Missing source file: {name}"); continue
        try:
            # Random cell access is used because the worksheets store each motor in a
            # column. Non-streaming mode avoids reparsing worksheet XML per cell.
            wb = load_workbook(path, data_only=True, read_only=False)
            if not wb.worksheets: errors.append(f"No worksheet in {name}"); continue
            ws = wb.worksheets[0]
            for col in range(4, ws.max_column + 1):
                motor_kind = clean(ws.cell(1, col).value)
                power = clean(ws.cell(2, col).value)
                frame = clean(ws.cell(4, col).value)
                if not power or not frame or not motor_kind: continue
                family = "BAX" if motor_kind.startswith("BAX") else ("BMX" if motor_kind == "BMX" else ("SMX" if motor_kind == "SMX" else None))
                if not family or (family_hint and family != family_hint): continue
                values = {}
                for row in range(1, ws.max_row + 1):
                    label = clean(ws.cell(row, 2).value)
                    sublabel = clean(ws.cell(row, 3).value)
                    cell=ws.cell(row, col)
                    value = excel_display(cell.value, cell.number_format)
                    if label:
                        key = label
                        if label in ("Bearing", "Material") and sublabel: key = f"{label} - {sublabel}"
                        values[key] = value
                    elif row in (21, 22):
                        values[f"Efficiency {clean(ws.cell(row, 3).value)}"] = value
                    elif row in (24, 25):
                        values[f"Power Factor {clean(ws.cell(row, 3).value)}"] = value
                eff = values.get("Efficiency Class") or class_hint
                if class_hint and eff != class_hint: continue
                brake_type = values.get("Brake type") or brake_hint
                records.append({"family": family, "efficiency": eff, "power": power, "frame": frame,
                    "brakeType": brake_type, "source": name, "column": col, "values": values})
            wb.close()
        except Exception as exc:
            errors.append(f"Could not read {name}: {exc}")
    return records, errors

RECORDS, SOURCE_ERRORS = read_sources()

def available():
    return {"families": sorted({r["family"] for r in RECORDS}), "classes": sorted({r["efficiency"] for r in RECORDS if r["efficiency"]}),
            "records": len(RECORDS), "sourceErrors": SOURCE_ERRORS}

def diagnostic_report():
    groups=defaultdict(list)
    for r in RECORDS: groups[(r["family"],r["efficiency"],r["power"])].append(r)
    issues=[]
    for (family,efficiency,power),rows in groups.items():
        configs=defaultdict(list)
        for r in rows: configs[r["brakeType"]].append(r)
        for brake_type,config_rows in configs.items():
            issue=conflicts(config_rows)
            if issue:
                issues.append({"family":family,"efficiency":efficiency,"power":power,"brakeType":brake_type,"fields":issue["fields"],"sources":issue["sources"],
                               "values":[{"source":r["source"],"values":{k:r["values"].get(k) for k in issue["fields"]}} for r in config_rows]})
    return {"recordCount":len(RECORDS),"sourceErrors":SOURCE_ERRORS,"conflicts":issues}

def select_options(q):
    family, cls, power = q.get("family"), q.get("efficiency"), q.get("power")
    base = [r for r in RECORDS if not family or r["family"] == family]
    powers = sorted({r["power"] for r in base}, key=lambda x: float(x))
    rows = [r for r in base if not power or r["power"] == power]
    poles = sorted({r["values"].get("No. of Poles") for r in rows if r["values"].get("No. of Poles")})
    if q.get("noOfPoles"):
        rows = [r for r in rows if r["values"].get("No. of Poles") == q["noOfPoles"]]
    efficiencies = sorted({r["efficiency"] for r in rows if r["efficiency"]})
    if cls:
        rows = [r for r in rows if r["efficiency"] == cls]

    def option_value(record, field):
        value = record["values"].get(field)
        if field == "Insulation Class" and value:
            match = re.search(r"\(([FH])\)", value, re.I)
            return match.group(1).upper() if match else None
        if field == "Degree of protection" and value:
            return re.sub(r"\s+", "", value).upper()
        if field == "Mounting" and value:
            return re.sub(r"^IM\s*", "", value, flags=re.I).upper()
        if field == "Method of cooling" and value:
            return re.sub(r"IC\s*(\d+)", r"IC \1", value, flags=re.I)
        return value

    config_fields = {
        "insulationClass": "Insulation Class",
        "protection": "Degree of protection",
        "mounting": "Mounting",
        "coolingMethod": "Method of cooling",
    }
    # Each option list is calculated from source rows matching all prior choices.
    option_rows = {}
    prior = list(rows)
    for query_key, source_key in config_fields.items():
        option_rows[query_key] = sorted({option_value(r, source_key) for r in prior if option_value(r, source_key)})
        selected = q.get(query_key)
        if selected:
            prior = [r for r in prior if option_value(r, source_key) == selected]
    brake_types = sorted({r["brakeType"] for r in prior if r["brakeType"]})
    selected_brake = q.get("brakeType")
    matches = [r for r in prior if not selected_brake or r["brakeType"] == selected_brake]
    frames = sorted({r["frame"] for r in matches})
    merged_sources = None
    if same_motor_data(matches):
        merged_sources = sorted({r["source"] for r in matches})
        matches = [matches[0]]
    record = matches[0] if len(matches) == 1 else None
    if record and merged_sources: record = dict(record, source=" + ".join(merged_sources))
    return {"powers": powers, "poles": poles, "efficiencies": efficiencies,
            "insulationClasses": option_rows["insulationClass"],
            "protections": option_rows["protection"], "mountings": option_rows["mounting"],
            "coolingMethods": option_rows["coolingMethod"], "brakeTypes": brake_types, "frames": frames, "record": record,
            "recordCount": len(matches), "conflict": conflicts(matches)}

def conflicts(rows):
    if len(rows) < 2: return None
    # AC and DC brake records are separate valid configurations. Only compare
    # duplicate records after the brake configuration has disambiguated them.
    if len({r.get("brakeType") for r in rows}) > 1: return None
    ignored = {"Motor Type", "Data Sheet No.", "PD Drg No.", "TB Drg No.", "Load Vs Efficiency curve", "Load Vs power factor curve", "Torque speed curve", "Current Vs time curve", "Current Vs Speed curve", "Thermal withstand curve", "Brake type", "Input supply voltage", "No. of Phase", "Brake coil voltage", "Brake torque", "Brake safety factor", "Brake reaction time", "Brake current"}
    keys = set().union(*(r["values"].keys() for r in rows)) - ignored
    diff = [k for k in sorted(keys) if len({r["values"].get(k) for r in rows}) > 1]
    return {"sources": [r["source"] for r in rows], "fields": diff} if diff else None

def same_motor_data(rows):
    return len(rows) > 1 and not conflicts(rows) and len({r["brakeType"] for r in rows}) == 1

def record_view(r):
    if not r: return None
    return {"family": r["family"], "efficiency": r["efficiency"], "power": r["power"], "frame": r["frame"], "brakeType": r["brakeType"], "source": r["source"], "values": r["values"]}

# Reference page rows and labels are retained. Only value glyphs are erased/redrawn.
PDF_ROWS = {
    "rated power": (2, 0), "number of poles": (3, 0), "frame size": (4, 0), "type of mounting": (5, 0),
    "rated voltage": (6, 0), "rated frequency": (8, 0), "stator connection": (7, 0), "rated current": (9, 0),
    "rated speed": (10, 0), "rated torque": (11, 0), "efficiency class": (12, 0), "duty type": (13, 0),
    "insulation class": (14, 0), "temperature rise class": (14, 0), "degree of protection": (15, 0),
    "method of cooling": (16, 0), "locked rotor current": (17, 0), "locked rotor torque": (18, 0),
    "breakdown torque": (19, 0), "reference standard": (52, 0), "bearing": (30, 1),
    "bearing life time": (32, 1), "lubrication type": (33, 1), "voltage variation": (38, 1),
    "frequency variation": (39, 1), "combined variation": (40, 1), "environmental condition": (26, 1),
    "direction of rotation": (None, 1), "housing material": (34, 1), "flange material": (35, 1),
    "motor weight": (36, 1), "lifting eyebot size": (37, 1), "terminal box position (view from nde)": (41, 1),
    "type of terminal box": (43, 1), "material of terminal box": (42, 1), "terminal size": (44, 1),
    "no. of terminals": (45, 1), "cable entry thread size": (46, 1), "no. of cable glands": (47, 1),
    "cable diameter": (48, 1), "max. cable size": (49, 1), "painting cycle": (50, 1),
    "color / paint shade **": (51, 1), "vibration severity grade": (53, 1),
    "sound pressure level - 50hz": (54, 1), "brake type": (55, 0), "brake torque": (59, 1),
    "brake safety factor": (60, 1), "input supply voltage": (56, 0), "no. of phase": (57, 0),
    "brake reaction time": (61, 1), "brake coil voltage": (58, 0), "brake current": (62, 1),
}

def pdf_value_for(label, rec):
    vals = rec["values"]
    row, side = PDF_ROWS.get(label, (None, None))
    if label == "rated power": return rec["power"]
    if label == "frame size": return rec["frame"]
    if label == "type of mounting":
        s=vals.get("Mounting"); return s.replace("IM ", "") if s else None
    if label == "bearing":
        de, nde = vals.get("Bearing - DE"), vals.get("Bearing - NDE")
        return f"DE {de} / NDE {nde}" if de and nde else de or nde
    if label == "voltage variation" or label == "frequency variation" or label == "combined variation":
        v=vals.get({"voltage variation":"Voltage variation","frequency variation":"Frequency variation","combined variation":"Combined variation"}[label])
        return v.replace("%", "").strip() if v else None
    if label == "sound pressure level - 50hz":
        return vals.get("Sound pressure level - 50Hz")
    if label == "input supply voltage": return vals.get("Input supply voltage")
    if label == "brake coil voltage": return vals.get("Brake coil voltage")
    if label == "brake type": return rec.get("brakeType")
    if label == "temperature rise class":
        v=vals.get("Insulation Class", "")
        m=re.search(r"utilized to\s+(\d+)\s*\(([A-Z])\)", v, re.I)
        return m.group(2) if m else None
    if label in ("locked rotor current", "locked rotor torque", "breakdown torque"):
        ratio_key={"locked rotor current":"Current - Locked rotor / Rated", "locked rotor torque":"Torque - Locked rotor / Rated", "breakdown torque":"Torque - Breakdown / Rated"}[label]
        base_key="Rated Current" if label=="locked rotor current" else "Rated Torque"
        ratio=vals.get(ratio_key); base=vals.get(base_key)
        # Excel supplies IA/IN and torque ratios, while the PDF cells are in A/Nm.
        # Convert only when both source operands are numeric; use the first current
        # entry because the reference's corresponding field is a single value.
        if not ratio or not base or ratio.lower()=="italy" or base.lower()=="italy": return None
        base_num=re.search(r"[-+]?\d+(?:\.\d+)?",base)
        ratio_num=re.search(r"[-+]?\d+(?:\.\d+)?",ratio)
        if not base_num or not ratio_num: return None
        try: return f"{Decimal(base_num.group())*Decimal(ratio_num.group()):.1f}"
        except InvalidOperation: return None
    if row is None: return None
    value = vals.get({"rated power":"Power Rating", "number of poles":"No. of Poles", "frame size":"Frame Size",
                      "type of mounting":"Mounting", "rated voltage":"Rated Voltage", "rated frequency":"Rated Frequency",
                      "stator connection":"Stator Connection", "rated current":"Rated Current", "rated speed":"Rated Speed",
                      "rated torque":"Rated Torque", "efficiency class":"Efficiency Class", "duty type":"Duty Type",
                      "insulation class":"Insulation Class", "degree of protection":"Degree of protection",
                      "method of cooling":"Method of cooling", "locked rotor current":"Current - Locked rotor / Rated",
                      "locked rotor torque":"Torque - Locked rotor / Rated", "breakdown torque":"Torque - Breakdown / Rated",
                      "reference standard":"Reference standard", "bearing life time":"Bearing life time",
                      "lubrication type":"Lubrication type", "voltage variation":"Voltage variation", "frequency variation":"Frequency variation",
                      "combined variation":"Combined variation", "environmental condition":"Environmental condition", "housing material":"Material - Housing",
                      "flange material":"Material - Flange", "motor weight":"Motor weight", "lifting eyebot size":"Lifting eyebolt size",
                      "terminal box position (view from nde)":"Terminal box position", "type of terminal box":"Type of terminal box",
                      "material of terminal box":"Material of terminal box", "terminal size":"Terminal size", "no. of terminals":"No of terminals",
                      "cable entry thread size":"Cable entry thread size", "no. of cable glands":"No. of cable glands",
                      "cable diameter":"Cable diameter (min - max)", "max. cable size":"Max. Cable size", "painting cycle":"Painting cycle",
                      "color / paint shade **":"Color / Paint Shade", "vibration severity grade":"Vibration severity grade",
                      "brake torque":"Brake torque", "brake safety factor":"Brake safety factor", "no. of phase":"No. of Phase",
                      "brake reaction time":"Brake reaction time", "brake current":"Brake current"}.get(label, ""))
    if value and value.lower() in {"italy", "n/a", "-"}: return None
    if label == "method of cooling" and value: return value.replace("IC411", "IC 411")
    if label == "degree of protection" and value: return value.replace("IP ", "IP")
    if label == "reference standard" and value: return value.replace(" : ", ":").replace(" / ", "/")
    if label == "insulation class" and value:
        m=re.search(r"\(([A-Z])\)", value)
        return m.group(1) if m else None
    return value

def generate_pdf(rec):
    if not PDF.exists(): raise ValueError("Reference PDF template is missing.")
    if not rec: raise ValueError("The selected motor combination is not available in the supplied workbooks.")
    doc=fitz.open(PDF); page=doc[0]
    logo_info=page.get_images(full=True)
    logo=None
    if logo_info:
        xref=logo_info[0][0]
        rects=page.get_image_rects(xref)
        if rects: logo=(fitz.Rect(rects[0]),doc.extract_image(xref)["image"])
    # Remove reference sample values while keeping the grid, labels, logo and page layout.
    spans=[]
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                spans.append(span)
    value_bands=[(220,298,65,498),(480,561,65,582)]
    for s in spans:
        x0,y0,x1,y1=s["bbox"]
        if any(x0>=a and x0<=b and c<=y0<=d for a,b,c,d in value_bands):
            page.add_redact_annot(fitz.Rect(x0-0.5,y0-0.6,x1+0.5,y1+0.7),fill=(1,1,1))
        if 220<=x0<=260 and 518<=y0<=582:
            page.add_redact_annot(fitz.Rect(x0-0.5,y0-0.6,x1+0.5,y1+0.7),fill=(1,1,1))
        if x0>=222 and y0>=32 and y1<=47 and x1<590:
            page.add_redact_annot(fitz.Rect(x0-0.5,y0-0.6,x1+0.5,y1+0.7),fill=(1,1,1))
        if not rec.get("brakeType") and s["text"].endswith("(Brake motor)"):
            page.add_redact_annot(fitz.Rect(x0-0.5,y0-0.6,x1+0.5,y1+0.7),fill=(1,1,1))
        if rec.get("brakeType")=="DC" and 275<=x0<=285 and 565<=y0<=572:
            # The reference's static suffix says AC; the DC workbook specifies a DC brake coil.
            page.add_redact_annot(fitz.Rect(x0-0.5,y0-0.6,x1+0.5,y1+0.7),fill=(1,1,1))
    page.apply_redactions(images=0,graphics=0)
    values={}
    for label,(row,side) in PDF_ROWS.items():
        if row is None or label=="temperature rise class": continue
        y={2:68.7,3:85.3,4:102.0,5:118.7,6:135.4,7:168.7,8:152.1,9:185.4,10:202.1,11:218.8,12:235.5,13:252.1,14:268.8,15:302.2,16:318.9,17:335.5,18:352.2,19:368.9,26:185.4,30:68.7,32:102.0,33:118.7,34:218.8,35:235.5,36:252.1,37:268.8,38:135.4,39:152.1,40:168.7,41:285.5,42:318.9,43:302.2,44:335.5,45:352.2,46:368.9,47:385.6,48:402.3,49:418.9,50:435.6,51:452.3,52:485.7,53:469.0,54:485.7,55:519.0,56:535.7,57:552.4,58:569.1,59:519.0,60:535.7,61:552.4,62:569.1}.get(row)
        if y is None: continue
        val=pdf_value_for(label,rec)
        if val:
            # The efficiency, power-factor rows are handled by their load-specific worksheet rows below.
            values[(side,y)]=val
    values[(0,68.7)]=rec["power"]
    temperature=pdf_value_for("temperature rise class",rec)
    if temperature: values[(0,285.5)]=temperature
    for y, key in [(385.6,"Efficiency 1"),(402.3,"Efficiency 0.75"),(418.9,"Efficiency 0.5"),
                   (435.6,"Power Factor 1"),(452.3,"Power Factor 0.75"),(469.0,"Power Factor 0.5")]:
        val=rec["values"].get(key)
        if val and val.lower()!="italy": values[(0,y)]=val
    leftx,rightx=224.2,487.9
    for (side,y),val in values.items():
        x=leftx if side==0 else rightx
        if len(val)>29: val=val[:27]+"…"
        max_width=(319-x if side==0 else 580-x)-4
        font_size=min(8.4,max(6.0,max_width/max(fitz.get_text_length(val,fontname="helv",fontsize=8.4),1)*8.4))
        page.insert_text((x,y+9.6),val,fontname="helv",fontsize=font_size,color=(0,0,0),overlay=True)
    # Record descriptor and revision metadata.
    desc=f"{rec['family']} {rec['frame']} {rec['values'].get('No. of Poles','')} {rec['values'].get('Rated Voltage','')} {rec['values'].get('Rated Frequency','')} {rec['values'].get('Mounting','')} P-{rec['power']}kW, {rec['efficiency']}, {rec['values'].get('Degree of protection','')}, {rec['values'].get('Duty Type','')}"
    page.insert_text((224.3,43.4),desc[:90],fontname="hebo",fontsize=8.2,color=(0,0,0),overlay=True)
    if logo: page.insert_image(logo[0],stream=logo[1],overlay=True)
    if not rec.get("brakeType"):
        page.insert_text((224.3,29.6),"Data sheet - Three phase - Squirrel cage motors",fontname="hebo",fontsize=10.06,color=(0,0,0),overlay=True)
    elif rec.get("brakeType")=="DC":
        page.insert_text((280.7,578.7),"(DC)   V",fontname="helv",fontsize=8.4,color=(0,0,0),overlay=True)
    return doc.tobytes(garbage=4,deflate=True)

class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args): print("%s - %s"%(self.address_string(),fmt%args))
    def respond(self, code, body, content_type="application/json; charset=utf-8", headers=None):
        self.send_response(code); self.send_header("Content-Type",content_type); self.send_header("Cache-Control","no-store")
        for k,v in (headers or {}).items(): self.send_header(k,v)
        self.end_headers(); self.wfile.write(body if isinstance(body,bytes) else body.encode())
    def do_GET(self):
        path=urlparse(self.path).path
        if path=="/api/meta": return self.respond(200,json.dumps(available()))
        if path=="/api/diagnostics": return self.respond(200,json.dumps(diagnostic_report()))
        if path=="/api/options":
            import urllib.parse
            return self.respond(200,json.dumps(select_options(dict(urllib.parse.parse_qsl(urlparse(self.path).query)))))
        if path=="/" or path=="/index.html": path="/index.html"
        if path.startswith("/assets/"): target=ROOT/path.lstrip("/")
        else: target=STATIC/path.lstrip("/")
        if not target.resolve().is_relative_to(ROOT.resolve()) or not target.is_file(): return self.respond(404,"Not found","text/plain")
        ctype=mimetypes.guess_type(target.name)[0] or "application/octet-stream"
        return self.respond(200,target.read_bytes(),ctype)
    def do_POST(self):
        if urlparse(self.path).path!="/api/datasheet": return self.respond(404,"Not found","text/plain")
        try:
            data=json.loads(self.rfile.read(int(self.headers.get("Content-Length",0))))
            required=("family","power","noOfPoles","efficiency","insulationClass","protection","mounting","coolingMethod")
            missing=[key for key in required if not data.get(key)]
            if missing: raise ValueError("Complete all motor configuration fields before generating the datasheet.")
            options=select_options(data); rec=options.get("record")
            if options.get("brakeTypes") and not data.get("brakeType"):
                raise ValueError("Select a brake configuration available in the source workbook.")
            if options.get("conflict"): raise ValueError("Conflicting workbook values found for this motor/power/class. Resolve the source data before generating a datasheet.")
            if options.get("recordCount")!=1 or not rec: raise ValueError("Selected combination is missing or ambiguous in the supplied Excel data.")
            pdf=generate_pdf(rec)
            filename=f"Motor_Datasheet_{rec['family']}_{rec['power']}kW_{rec['efficiency']}.pdf".replace(" ","")
            return self.respond(200,pdf,"application/pdf",{"Content-Disposition":f'attachment; filename="{filename}"'})
        except Exception as exc: return self.respond(400,json.dumps({"error":str(exc)}))

if __name__=="__main__":
    host="127.0.0.1"; port=int(sys.argv[1]) if len(sys.argv)>1 else 8765
    print(f"MGM Varvel selector available at http://{host}:{port}")
    if SOURCE_ERRORS: print("Source warnings:", *SOURCE_ERRORS, sep="\n- ")
    ThreadingHTTPServer((host,port),Handler).serve_forever()
